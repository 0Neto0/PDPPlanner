from datetime import datetime

from openpyxl import Workbook

from pdpplanner.excel_readers import (
    read_production_operations,
    read_production_order_list,
    read_sap_orders,
    read_staff_calendar,
)


def _make_sap_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2024"
    sheet["B5"] = "PLAZO DE ENTREGA"
    sheet["D5"] = "FECHA DE ENTREGA"
    sheet["E5"] = "Nº PEDIDO SAP"
    sheet["G5"] = "DESCRIPCION"
    sheet["H5"] = "O.T."
    sheet["J5"] = "ESTADO"
    sheet["C6"] = datetime(2026, 1, 2)  # Column C must be ignored.
    sheet["B6"] = datetime(2026, 1, 10)
    sheet["D6"] = datetime(2026, 1, 9)
    sheet["E6"] = 4500000001
    sheet["G6"] = "MATERIAL"
    sheet["H6"] = 2300001
    sheet["J6"] = "rECIBIDO"
    sheet["E7"] = 4500000002
    sheet["G7"] = "MATERIAL"
    sheet["H7"] = 2300002
    sheet["J7"] = "ESTADO NUEVO"
    sheet["E8"] = 4500000002
    sheet["H8"] = 2300003
    sheet["J8"] = "EN TRAMITE"
    sheet["E9"] = 4500000003
    sheet["G9"] = "MATERIAL SIN OT"
    sheet["J9"] = "ESTADO NUEVO"
    workbook.save(path)


def test_read_sap_orders_preserves_rows_and_reports_issues(tmp_path):
    path = tmp_path / "sap.xlsx"
    _make_sap_workbook(path)
    result = read_sap_orders(path)
    assert len(result.records) == 4
    assert result.records[0].ot == "2300001"
    assert result.records[0].order_id == "4500000001"
    assert result.records[0].expected_date == datetime(2026, 1, 10)
    assert result.records[0].actual_date == datetime(2026, 1, 9)
    assert {issue.code for issue in result.issues} >= {
        "SAP_MISSING_OT", "SAP_UNKNOWN_STATE", "SAP_ORDER_ID_MULTIPLE_OTS",
    }


def test_missing_sheet_is_reported_as_error(tmp_path):
    path = tmp_path / "sap.xlsx"
    Workbook().save(path)
    result = read_sap_orders(path)
    assert result.records == []
    assert result.has_errors
    assert result.issues[0].code == "SAP_SHEET_MISSING"


def test_read_production_operations_handles_variable_ot_block_lengths(tmp_path):
    path = tmp_path / "production.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PdP PRODUCCIÓN"
    sheet["B19"] = "Nº OT"
    sheet["C19"] = "TIPO MÁQUINA"
    sheet["F19"] = "H. PRESUP."
    sheet["B29"], sheet["C29"], sheet["E29"], sheet["F29"] = 1001, "Conjunto A", "CALDERERÍA", 12
    sheet["C30"], sheet["E30"], sheet["F30"] = "Conjunto A", "TORNO", 4
    sheet["B31"], sheet["C31"], sheet["E31"], sheet["F31"] = 1002, "Conjunto B", "FRESA", 7
    workbook.save(path)
    result = read_production_operations(path)
    assert [(item.ot, item.machine, item.budget_hours) for item in result.records] == [
        ("1001", "CALDERERÍA", 12), ("1001", "TORNO", 4), ("1002", "FRESA", 7),
    ]
    assert not result.has_errors


def test_production_operation_without_ot_is_reported(tmp_path):
    path = tmp_path / "production.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PdP PRODUCCIÓN"
    sheet["B19"], sheet["C19"], sheet["F19"] = "Nº OT", "TIPO MÁQUINA", "H. PRESUP."
    sheet["E29"] = "TORNO"
    workbook.save(path)
    result = read_production_operations(path)
    assert result.records == []
    assert result.issues[0].code == "PRODUCTION_OPERATION_WITHOUT_OT"


def test_missing_production_sheet_is_reported_as_error(tmp_path):
    path = tmp_path / "production.xlsx"
    Workbook().save(path)
    result = read_production_operations(path)
    assert result.records == []
    assert result.has_errors
    assert result.issues[0].code == "PRODUCTION_SHEET_MISSING"


def test_read_production_order_list_maps_columns(tmp_path):
    path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Listado Pedidos"
    sheet["B2"], sheet["G2"], sheet["H2"], sheet["N2"] = "Nº OT", "Producto", "CLIENTE", "Estado"
    sheet["A3"], sheet["B3"], sheet["F3"] = 2, 2500572, "SI"
    sheet["G3"], sheet["H3"], sheet["N3"] = "MESA RODILLOS", "GOSIMAT", "EN PROCESO"
    sheet["I3"], sheet["J3"], sheet["K3"], sheet["L3"] = 1, 0.5, 1, 0.75
    workbook.save(path)
    result = read_production_order_list(path)
    assert len(result.records) == 1
    item = result.records[0]
    assert (item.ot, item.product, item.client, item.status) == (
        "2500572", "MESA RODILLOS", "GOSIMAT", "EN PROCESO",
    )
    assert item.progress == (1, 0.5, 1, 0.75)
    assert not result.has_errors


def test_read_staff_calendar_captures_marked_unavailable_dates(tmp_path):
    path = tmp_path / "staff.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vacaciones-Bajas"
    sheet["B13"], sheet["C13"], sheet["D13"] = "OPERARIO", "SECCION", "PUESTO"
    sheet["J12"], sheet["K12"] = datetime(2026, 1, 1), datetime(2026, 1, 2)
    sheet["B15"], sheet["C15"], sheet["D15"] = "Operario A", "TORNO", "Tornero"
    sheet["J15"] = 1
    workbook.save(path)
    result = read_staff_calendar(path)
    assert len(result.records) == 1
    assert result.records[0].name == "Operario A"
    assert result.records[0].section == "TORNO"
    assert result.records[0].unavailable_dates == (datetime(2026, 1, 1),)
    assert not result.has_errors

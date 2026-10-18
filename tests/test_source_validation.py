from datetime import datetime

from openpyxl import Workbook

from pdpplanner.source_validation import read_and_validate_sources


def test_read_and_validate_sources_returns_counts_and_findings(tmp_path):
    management_path = tmp_path / "management.xlsx"
    management = Workbook()
    production = management.active
    production.title = "PdP PRODUCCIÓN"
    production["B19"], production["C19"], production["F19"] = "Nº OT", "TIPO MÁQUINA", "H. PRESUP."
    production["B29"], production["C29"], production["E29"], production["F29"] = 2300001, "Equipo A", "TORNO", 8

    order_list = management.create_sheet("Listado Pedidos")
    order_list["B2"], order_list["G2"], order_list["H2"], order_list["N2"] = "Nº OT", "Producto", "CLIENTE", "Estado"
    order_list["B3"], order_list["G3"], order_list["H3"], order_list["N3"] = 2300001, "Equipo A", "Cliente A", "EN PROCESO"

    staff = management.create_sheet("Vacaciones-Bajas")
    staff["J12"] = datetime(2026, 1, 1)
    staff["B13"], staff["C13"], staff["D13"] = "OPERARIO", "SECCION", "PUESTO"
    staff["B15"], staff["C15"], staff["D15"] = "Operario A", "TORNO", "Tornero"
    staff["J15"] = 1
    management.save(management_path)

    sap_path = tmp_path / "sap.xlsx"
    sap_book = Workbook()
    sap = sap_book.active
    sap.title = "2024"
    sap["B5"], sap["D5"], sap["E5"] = "PLAZO DE ENTREGA", "FECHA DE ENTREGA", "Nº PEDIDO SAP"
    sap["G5"], sap["H5"], sap["J5"] = "DESCRIPCION", "O.T.", "ESTADO"
    sap["B6"], sap["D6"], sap["E6"] = datetime(2026, 1, 5), None, 4500000001
    sap["G6"], sap["H6"], sap["J6"] = "Material", 2300001, "EN TRAMITE"
    sap_book.save(sap_path)

    bundle = read_and_validate_sources(management_path, sap_path)
    report = bundle.report()

    assert report["summary"]["record_counts"] == {
        "production_operations": 1,
        "production_orders": 1,
        "staff_calendar": 1,
        "sap_orders": 1,
    }
    assert report["summary"]["issue_count"] == 0
    assert report["summary"]["has_errors"] is False
    assert report["issues"] == []


def test_missing_source_sheets_are_aggregated_as_errors(tmp_path):
    management_path = tmp_path / "management.xlsx"
    Workbook().save(management_path)
    sap_path = tmp_path / "sap.xlsx"
    Workbook().save(sap_path)

    bundle = read_and_validate_sources(management_path, sap_path)

    assert bundle.has_errors
    assert {issue.code for issue in bundle.issues} >= {
        "PRODUCTION_SHEET_MISSING",
        "ORDER_LIST_SHEET_MISSING",
        "STAFF_SHEET_MISSING",
        "SAP_SHEET_MISSING",
    }
    assert bundle.summary()["issue_count"] == 4

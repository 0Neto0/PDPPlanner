"""Read-only Excel adapters. These functions never save or modify source workbooks."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from pdpplanner.models import ProductionOperation, ProductionOrderInfo, ReadResult, SapOrder, StaffMember, ValidationIssue
from pdpplanner.order_status import classify_state, normalize_ot

SAP_SHEET = "2024"
SAP_HEADER_ROW = 5
PRODUCTION_SHEET = "PdP PRODUCCIÓN"
PRODUCTION_HEADER_ROW = 19
PRODUCTION_FIRST_DATA_ROW = 29
ORDER_LIST_SHEET = "Listado Pedidos"
ORDER_LIST_HEADER_ROW = 2
STAFF_SHEET = "Vacaciones-Bajas"
STAFF_HEADER_ROW = 13
STAFF_FIRST_DATA_ROW = 15
SAP_COLUMNS = {"expected_date": 2, "actual_date": 4, "order_id": 5, "description": 7, "ot": 8, "state": 10}
EXPECTED_HEADERS = {"expected_date": "PLAZO DE ENTREGA", "actual_date": "FECHA DE ENTREGA", "order_id": "Nº PEDIDO SAP", "description": "DESCRIPCION", "ot": "O.T.", "state": "ESTADO"}


def _clean_identifier(value: Any) -> str | None:
    if value is None or not str(value).strip():
        return None
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().upper().split())


def read_sap_orders(path: str | Path, *, sheet_name: str = SAP_SHEET, header_row: int = 5) -> ReadResult[SapOrder]:
    """Read SAP rows and report ambiguous data without inferring missing links."""
    source, records, issues = str(path), [], []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return ReadResult([], [ValidationIssue("SAP_SHEET_MISSING", f"No existe la hoja requerida {sheet_name!r}.", source, severity="error")])
        worksheet = workbook[sheet_name]
        header_values = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, max_col=max(SAP_COLUMNS.values()), values_only=True), ())
        for field, column in SAP_COLUMNS.items():
            actual = header_values[column - 1] if len(header_values) >= column else None
            actual_text = _normalize_header(actual)
            if field == "order_id":
                actual_text = actual_text.replace("N°", "Nº").replace("NO.", "Nº")
            if actual_text != _normalize_header(EXPECTED_HEADERS[field]):
                issues.append(ValidationIssue("SAP_HEADER_MISMATCH", f"Encabezado inesperado en columna {column}, fila {header_row}: {actual!r}; se esperaba {EXPECTED_HEADERS[field]!r}.", source, row=header_row))
        for row_number, values in enumerate(worksheet.iter_rows(min_row=header_row + 1, max_col=max(SAP_COLUMNS.values()), values_only=True), start=header_row + 1):
            raw = {field: values[column - 1] if len(values) >= column else None for field, column in SAP_COLUMNS.items()}
            if all(value is None or not str(value).strip() for value in raw.values()):
                continue
            order_id, raw_ot = _clean_identifier(raw["order_id"]), _clean_identifier(raw["ot"])
            ot, state = normalize_ot(raw_ot) if raw_ot else None, str(raw["state"] or "").strip()
            records.append(SapOrder(row_number, order_id, ot, state, raw["expected_date"], raw["actual_date"], str(raw["description"] or "").strip()))
            if not ot:
                issues.append(ValidationIssue("SAP_MISSING_OT", "El registro SAP no tiene OT; no se asociará por inferencia.", source, row=row_number))
            if not order_id:
                issues.append(ValidationIssue("SAP_MISSING_ORDER_ID", "El registro SAP no tiene número de pedido.", source, row=row_number))
            if classify_state(state) == "unknown":
                issues.append(ValidationIssue("SAP_UNKNOWN_STATE", f"Estado SAP no reconocido: {state!r}.", source, row=row_number))
        order_rows: dict[str, list[SapOrder]] = defaultdict(list)
        for record in records:
            if record.order_id:
                order_rows[record.order_id].append(record)
        for order_id, duplicates in order_rows.items():
            ots = {record.ot for record in duplicates if record.ot}
            if len(ots) > 1:
                row_list = ", ".join(str(record.source_row) for record in duplicates)
                issues.append(ValidationIssue("SAP_ORDER_ID_MULTIPLE_OTS", f"El pedido SAP {order_id!r} aparece asociado a varias OT ({', '.join(sorted(ots))}); filas: {row_list}. No se fusiona.", source, row=duplicates[0].source_row))
    finally:
        workbook.close()
    return ReadResult(records, issues)


def read_production_operations(path: str | Path, *, sheet_name: str = PRODUCTION_SHEET, header_row: int = PRODUCTION_HEADER_ROW, first_data_row: int = PRODUCTION_FIRST_DATA_ROW) -> ReadResult[ProductionOperation]:
    """Read production operations without assuming 8 rows per OT."""
    source, records, issues = str(path), [], []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return ReadResult([], [ValidationIssue("PRODUCTION_SHEET_MISSING", f"No existe la hoja requerida {sheet_name!r}.", source, severity="error")])
        worksheet = workbook[sheet_name]
        headers = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, min_col=2, max_col=6, values_only=True), ())
        expected = {0: "Nº OT", 1: "TIPO MÁQUINA", 4: "H. PRESUP."}
        for index, expected_header in expected.items():
            actual = headers[index] if len(headers) > index else None
            if _normalize_header(actual).replace("N°", "Nº") != _normalize_header(expected_header):
                issues.append(ValidationIssue("PRODUCTION_HEADER_MISMATCH", f"Encabezado inesperado en columna {index + 2}, fila {header_row}: {actual!r}; se esperaba {expected_header!r}.", source, row=header_row))
        active_ot = None
        for row_number, values in enumerate(worksheet.iter_rows(min_row=first_data_row, min_col=2, max_col=6, values_only=True), start=first_data_row):
            raw_ot, description, client, machine, budget_hours = (list(values) + [None] * 5)[:5]
            machine_text = str(machine or "").strip()
            if raw_ot is not None and str(raw_ot).strip():
                raw_ot_text = str(raw_ot).strip()
                if raw_ot_text.upper() not in {"OT", "Nº OT", "N° OT"}:
                    active_ot = normalize_ot(raw_ot_text)
            if not machine_text:
                if not any(value is not None and str(value).strip() for value in values):
                    active_ot = None
                continue
            if not active_ot:
                issues.append(ValidationIssue("PRODUCTION_OPERATION_WITHOUT_OT", "Hay una operación con máquina pero sin OT identificable.", source, row=row_number))
                continue
            records.append(ProductionOperation(row_number, active_ot, str(description or "").strip(), str(client or "").strip(), machine_text, budget_hours))
    finally:
        workbook.close()
    return ReadResult(records, issues)


def read_production_order_list(path: str | Path, *, sheet_name: str = ORDER_LIST_SHEET, header_row: int = ORDER_LIST_HEADER_ROW) -> ReadResult[ProductionOrderInfo]:
    """Read order metadata from ``Listado Pedidos``; formulas are not recalculated."""
    source, records, issues = str(path), [], []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return ReadResult([], [ValidationIssue("ORDER_LIST_SHEET_MISSING", f"No existe la hoja requerida {sheet_name!r}.", source, severity="error")])
        worksheet = workbook[sheet_name]
        headers = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, max_col=14, values_only=True), ())
        expected = {2: "Nº OT", 7: "PRODUCTO", 8: "CLIENTE", 14: "ESTADO"}
        for column, expected_header in expected.items():
            actual = headers[column - 1] if len(headers) >= column else None
            if _normalize_header(actual).replace("N°", "Nº") != _normalize_header(expected_header):
                issues.append(ValidationIssue("ORDER_LIST_HEADER_MISMATCH", f"Encabezado inesperado en columna {column}, fila {header_row}: {actual!r}; se esperaba {expected_header!r}.", source, row=header_row))
        for row_number, values in enumerate(worksheet.iter_rows(min_row=header_row + 1, max_col=14, values_only=True), start=header_row + 1):
            row = list(values) + [None] * 14
            if not any(value is not None and str(value).strip() for value in row[:14]):
                continue
            raw_ot, product = _clean_identifier(row[1]), str(row[6] or "").strip()
            if not raw_ot:
                if product:
                    issues.append(ValidationIssue("ORDER_LIST_MISSING_OT", "La fila de listado de producción no tiene OT.", source, row=row_number))
                continue
            records.append(ProductionOrderInfo(row_number, normalize_ot(raw_ot), row[0], str(row[5] or "").strip(), product, str(row[7] or "").strip(), str(row[13] or "").strip(), tuple(row[8:12])))
    finally:
        workbook.close()
    return ReadResult(records, issues)


def read_staff_calendar(path: str | Path, *, sheet_name: str = STAFF_SHEET, header_row: int = STAFF_HEADER_ROW, first_data_row: int = STAFF_FIRST_DATA_ROW) -> ReadResult[StaffMember]:
    """Read staff identity and explicitly marked unavailable calendar dates."""
    source, records, issues = str(path), [], []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return ReadResult([], [ValidationIssue("STAFF_SHEET_MISSING", f"No existe la hoja requerida {sheet_name!r}.", source, severity="error")])
        worksheet = workbook[sheet_name]
        date_row, first_calendar_column = 12, 10
        date_values = next(worksheet.iter_rows(min_row=date_row, max_row=date_row, min_col=first_calendar_column, max_col=worksheet.max_column, values_only=True), ())
        dates = {column: value for column, value in enumerate(date_values, start=first_calendar_column) if isinstance(value, (date, datetime))}
        if not dates:
            issues.append(ValidationIssue("STAFF_CALENDAR_DATES_MISSING", "No se pudieron leer fechas del calendario (fila 12, desde J).", source, row=date_row))
        for row_number, values in enumerate(worksheet.iter_rows(min_row=first_data_row, max_row=worksheet.max_row, min_col=2, max_col=worksheet.max_column, values_only=True), start=first_data_row):
            row = list(values)
            name, section, role = row[0], row[1], row[2]
            if not str(name or "").strip():
                continue
            unavailable = []
            for column, calendar_date in dates.items():
                index = column - 2
                if index < len(row) and row[index] is not None and str(row[index]).strip():
                    unavailable.append(calendar_date)
            records.append(StaffMember(row_number, str(name).strip(), str(section or "").strip(), str(role or "").strip(), tuple(unavailable)))
    finally:
        workbook.close()
    return ReadResult(records, issues)
